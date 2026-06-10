# -*- coding: utf-8 -*-
from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMS = ROOT / "diagrams"
RENDERED = DIAGRAMS / "rendered"
OUTPUT = DOCS / "TCG.Bet-Documentacao-Final.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
TEXT = RGBColor(30, 30, 30)
MUTED = RGBColor(90, 90, 90)
TABLE_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def font(size=28, bold=False):
    candidates = [
        r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\calibrib.ttf" if bold else r"C:\Windows\Fonts\calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def ensure_dirs():
    DOCS.mkdir(exist_ok=True)
    DIAGRAMS.mkdir(exist_ok=True)
    RENDERED.mkdir(exist_ok=True)


def draw_box(draw, xy, text, fill="#F7FAFC", outline="#2E74B5", text_fill="#0B2545", width=3):
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=width)
    x1, y1, x2, y2 = xy
    f = font(25, True)
    lines = wrap_text(text, f, x2 - x1 - 28)
    total_h = len(lines) * 30
    y = y1 + ((y2 - y1) - total_h) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=f)
        draw.text((x1 + ((x2 - x1) - (bbox[2] - bbox[0])) / 2, y), line, font=f, fill=text_fill)
        y += 30


def wrap_text(text, fnt, max_width):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if ImageDraw.Draw(Image.new("RGB", (1, 1))).textlength(candidate, font=fnt) <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def arrow(draw, start, end, color="#444444", width=4):
    draw.line([start, end], fill=color, width=width)
    ex, ey = end
    sx, sy = start
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        points = [(ex, ey), (ex - 16 * direction, ey - 9), (ex - 16 * direction, ey + 9)]
    else:
        direction = 1 if ey > sy else -1
        points = [(ex, ey), (ex - 9, ey - 16 * direction), (ex + 9, ey - 16 * direction)]
    draw.polygon(points, fill=color)


def create_diagram_images():
    # Architecture diagram
    img = Image.new("RGB", (1600, 900), "white")
    d = ImageDraw.Draw(img)
    title_font = font(34, True)
    d.text((60, 38), "Arquitetura geral - monolito modular em camadas", font=title_font, fill="#0B2545")
    boxes = {
        "Browser / Public": (80, 150, 360, 250),
        "Express Routes": (460, 150, 740, 250),
        "Services": (840, 150, 1120, 250),
        "Application\nUse Cases": (460, 360, 740, 480),
        "Domain Rules": (80, 360, 360, 480),
        "Application\nPorts": (840, 360, 1120, 480),
        "Infrastructure\nAdapters": (460, 610, 740, 730),
        "OracleDB / JWT /\nNodemailer": (840, 610, 1120, 730),
    }
    for label, xy in boxes.items():
        draw_box(d, xy, label, fill="#F7FAFC" if "Application" not in label else "#E8EEF5")
    arrow(d, (360, 200), (460, 200))
    arrow(d, (740, 200), (840, 200))
    arrow(d, (980, 250), (980, 360))
    arrow(d, (840, 420), (740, 420))
    arrow(d, (460, 420), (360, 420))
    arrow(d, (600, 480), (600, 610))
    arrow(d, (740, 670), (840, 670))
    d.text((80, 800), "Fonte versionada: diagrams/architecture.mmd", font=font(22), fill="#555555")
    img.save(RENDERED / "architecture.png")

    # GoF class diagram
    img = Image.new("RGB", (1600, 1000), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "Classes centrais dos padroes GoF", font=title_font, fill="#0B2545")
    draw_box(d, (80, 150, 390, 270), "PayoutStrategy\n<<interface>>", fill="#F4F6F9")
    draw_box(d, (80, 390, 390, 510), "PoolPayoutStrategy\nStrategy concreta", fill="#E8EEF5")
    draw_box(d, (520, 250, 900, 390), "FinishEventUseCase\nusa Strategy e Factory", fill="#F7FAFC")
    draw_box(d, (1030, 150, 1410, 290), "WalletTransactionFactory\nFactory", fill="#E8EEF5")
    draw_box(d, (1030, 430, 1410, 570), "WalletRepository\n<<interface>>", fill="#F4F6F9")
    draw_box(d, (1030, 700, 1410, 840), "OracleWalletRepository\nAdapter", fill="#E8EEF5")
    arrow(d, (235, 390), (235, 270))
    arrow(d, (520, 320), (390, 210))
    arrow(d, (900, 320), (1030, 220))
    arrow(d, (900, 365), (1030, 500))
    arrow(d, (1220, 700), (1220, 570))
    d.text((80, 915), "Fonte versionada: diagrams/gof-class-diagram.mmd", font=font(22), fill="#555555")
    img.save(RENDERED / "gof-class-diagram.png")

    # Sequence diagram
    img = Image.new("RGB", (1700, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((60, 38), "Fluxo central - realizar aposta", font=title_font, fill="#0B2545")
    participants = ["Usuario", "Route", "Service", "UseCase", "Auth", "Wallets", "Events", "Bets", "OracleDB"]
    xs = [90, 260, 430, 620, 820, 1010, 1190, 1360, 1530]
    for x, name in zip(xs, participants):
        draw_box(d, (x - 70, 130, x + 70, 200), name, fill="#F7FAFC")
        d.line([(x, 210), (x, 950)], fill="#B8C2CC", width=2)
    messages = [
        (0, 1, "token + aposta"),
        (1, 2, "betEvent"),
        (2, 3, "execute"),
        (3, 4, "authenticate"),
        (4, 8, "valida JWT"),
        (3, 6, "findById"),
        (3, 5, "findByOwnerId"),
        (3, 5, "setBalance"),
        (3, 7, "create bet"),
        (3, 5, "addHistory"),
        (2, 8, "commit"),
        (1, 0, "201 criado"),
    ]
    y = 250
    small = font(19)
    for src, dst, label in messages:
        start = (xs[src], y)
        end = (xs[dst], y)
        arrow(d, start, end, width=3)
        mid = (start[0] + end[0]) / 2
        d.text((mid - 55, y - 28), label, font=small, fill="#333333")
        y += 58
    d.text((80, 1010), "Fonte versionada: diagrams/place-bet-sequence.mmd", font=font(22), fill="#555555")
    img.save(RENDERED / "place-bet-sequence.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                set_cell_width(row.cells[idx], width)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    code = doc.styles.add_style("CodeBlock", 1)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor(40, 40, 40)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.text = "TCG.Bet - Documentacao de Projeto de Software"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = MUTED


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TCG.Bet")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Documentacao Final - Projeto de Software com Arquitetura Justificada")
    run.font.size = Pt(14)
    run.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Grupo: preencher nomes dos integrantes | Repositorio: https://github.com/KachanChK/TCG.Bet")
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED

    add_callout(
        doc,
        "Resumo executivo",
        "O TCG.Bet e uma aplicacao web de apostas em eventos futuros. A melhoria arquitetural adotou monolito modular, camadas inspiradas em Clean Architecture/Hexagonal, ADRs, OpenAPI, testes automatizados e padroes GoF aplicados a problemas reais do dominio."
    )


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(text)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_code(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    p = cell.paragraphs[0]
    for line in dedent(code).strip().splitlines():
        r = p.add_run(line)
        r.font.name = "Consolas"
        r.font.size = Pt(8)
        p.add_run("\n")
    doc.add_paragraph()


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.cell(0, idx)
        set_cell_shading(cell, TABLE_FILL)
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.color.rgb = DARK_BLUE
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def add_image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.4))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].font.size = Pt(9)
    cap.runs[0].font.color.rgb = MUTED


def build_doc():
    ensure_dirs()
    create_diagram_images()

    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("Sumario", level=1)
    add_numbered(doc, [
        "Introducao",
        "Atributos de qualidade e decisoes arquiteturais",
        "Estilo arquitetural",
        "Aplicacao dos principios SOLID",
        "Clean Code",
        "Padroes de projeto GoF",
        "Design de API",
        "Diagramas e modelos",
        "Conclusoes",
        "Referencias bibliograficas",
        "Anexos"
    ])

    doc.add_heading("1. Introducao", level=1)
    doc.add_paragraph(
        "O TCG.Bet e um software web para gerenciamento de apostas em eventos futuros. O sistema atende usuarios que desejam criar eventos, apostar, movimentar carteira e acompanhar historico financeiro, alem de moderadores responsaveis por avaliar eventos e finalizar resultados."
    )
    doc.add_paragraph(
        "A escolha do tema e adequada ao trabalho porque envolve regras de negocio nao triviais: autenticacao, autorizacao por papel, operacoes financeiras, historico de carteira, moderacao e calculo de premiacao. Esses pontos permitem demonstrar arquitetura, atributos de qualidade, SOLID, Clean Code e padroes GoF com evidencias objetivas no codigo."
    )
    doc.add_heading("Objetivo geral", level=2)
    doc.add_paragraph(
        "Desenvolver e documentar uma aplicacao de apostas em eventos futuros com arquitetura justificada, contrato REST formal e melhoria estrutural do codigo para favorecer manutencao e testabilidade."
    )
    doc.add_heading("Objetivos especificos", level=2)
    add_bullets(doc, [
        "Permitir cadastro, login e autenticacao de usuarios por JWT.",
        "Permitir criacao, avaliacao, cancelamento e finalizacao de eventos.",
        "Permitir deposito, saque, aposta e registro de historico financeiro.",
        "Separar regras de negocio de detalhes de infraestrutura.",
        "Documentar decisoes arquiteturais por ADRs e especificar a API com OpenAPI.",
        "Demonstrar SOLID, Clean Code e tres padroes GoF aplicados ao projeto."
    ])
    doc.add_heading("Problema e publico-alvo", level=2)
    doc.add_paragraph(
        "O problema tratado e a ausencia de uma plataforma simples para concentrar eventos de aposta, validacao por moderadores, participacao dos usuarios e controle de saldo. O publico-alvo sao usuarios adultos interessados em apostas sobre eventos futuros e moderadores que controlam a qualidade e o encerramento dos eventos."
    )

    doc.add_heading("2. Atributos de Qualidade e Decisoes Arquiteturais", level=1)
    doc.add_heading("2A. Atributos de qualidade prioritarios", level=2)
    add_table(doc, ["Atributo", "Justificativa", "Decisoes arquiteturais", "Metricas observaveis"], [
        [
            "Seguranca",
            "O sistema lida com autenticacao, autorizacao e operacoes financeiras. Um erro permitiria apostas indevidas ou uso de rotas de moderador por usuarios comuns.",
            "JWT para autenticacao, AuthPort/JwtOracleAuthAdapter, validacao de papel em casos de uso, separacao entre usuario e moderador e hash de senha com scrypt.",
            "100% das rotas privadas exigem Bearer token; rotas de moderador validam isModerator; tokens com expiracao de 8h; senhas novas nao sao persistidas em texto puro."
        ],
        [
            "Confiabilidade e consistencia",
            "Apostas, saldo e historico financeiro precisam permanecer coerentes, principalmente quando uma aposta debita saldo e registra historico.",
            "Operacoes financeiras executadas em uma conexao Oracle e confirmadas com commit apenas em sucesso; valores monetarios convertidos para centavos.",
            "Build sem erros; testes de dominio passando; nenhuma operacao de carteira deve aceitar valor menor ou igual a zero."
        ],
        [
            "Manutenibilidade e testabilidade",
            "O projeto precisa evoluir e ser defendido academicamente. Regras misturadas com SQL e Express dificultavam testes e mudancas.",
            "Camadas domain/application/infrastructure, portas pequenas, casos de uso focados, Strategy para premiacao e Factory para transacoes.",
            "Casos de uso centrais dependem de interfaces; 9 testes automatizados de dominio executados por npm test."
        ],
    ], widths=[1800, 2700, 3000, 1860])

    doc.add_heading("2B. Registro de Decisoes Arquiteturais", level=2)
    doc.add_paragraph("Os ADRs ficam versionados na pasta /adrs do repositorio. O ADR 0005 registra uma decisao modificada: a primeira versao usava servicos com SQL inline; a refatoracao moveu fluxos centrais para casos de uso e portas.")
    add_table(doc, ["ADR", "Titulo", "Status", "Resumo"], [
        ["0001", "Adotar monolito modular", "Accepted", "Mantem um unico deploy com separacao modular."],
        ["0002", "Organizar internamente por camadas", "Accepted", "Separa domain, application, infrastructure e HTTP."],
        ["0003", "Especificar API REST com OpenAPI", "Accepted", "Formaliza contrato REST em docs/openapi.yaml."],
        ["0004", "Aplicar Strategy, Factory e Adapter", "Accepted", "Usa GoF onde ha problemas reais de dominio."],
        ["0005", "Revisar servicos com SQL inline", "Accepted", "Modifica decisao inicial e melhora testabilidade."],
    ], widths=[900, 2800, 1200, 4460])

    doc.add_heading("3. Estilo Arquitetural", level=1)
    doc.add_heading("3A. Plano macro", level=2)
    doc.add_paragraph(
        "O sistema foi mantido como monolito modular. A escolha e coerente com o tamanho do projeto, com a maturidade esperada de um trabalho academico e com a necessidade de consistencia nas operacoes financeiras. Microsservicos aumentariam custo operacional, exigiriam comunicacao distribuida e complicariam transacoes sem trazer beneficio proporcional ao escopo atual."
    )
    doc.add_paragraph(
        "O monolito modular favorece seguranca e confiabilidade porque mantem autorizacao, apostas e carteira sob o mesmo processo e banco. A modularidade interna preserva manutenibilidade sem introduzir deploys independentes prematuros."
    )
    doc.add_heading("3B. Plano interno", level=2)
    doc.add_paragraph(
        "A organizacao interna e inspirada em Clean Architecture e Arquitetura Hexagonal. As regras puras ficam em domain; os casos de uso e portas ficam em application; os detalhes de OracleDB e JWT ficam em infrastructure; routes e services fazem a entrada HTTP e preservam compatibilidade com a API existente."
    )
    add_image(doc, RENDERED / "architecture.png", "Figura 1 - Arquitetura geral. Fonte: diagrams/architecture.mmd.")

    doc.add_heading("4. Aplicacao dos Principios SOLID", level=1)
    solid_rows = [
        ("Single Responsibility", "Cada classe deve ter uma unica razao para mudar.", "MoveWalletFundsUseCase concentra a regra de deposito/saque; WalletTransactionFactory cria transacoes; OracleWalletRepository persiste dados.", "Reduz mistura entre regra, SQL e HTTP."),
        ("Open-Closed", "Entidades devem estar abertas para extensao e fechadas para modificacao.", "FinishEventUseCase depende de PayoutStrategy.", "Um novo algoritmo de premiacao pode ser criado sem alterar o caso de uso."),
        ("Liskov Substitution", "Implementacoes de uma abstracao devem poder substituir a abstracao sem quebrar o cliente.", "OracleWalletRepository implementa WalletRepository.", "O caso de uso pode usar repositorio Oracle ou fake de teste."),
        ("Interface Segregation", "Clientes nao devem depender de metodos que nao usam.", "WalletRepository, BetRepository, EventRepository e AuthPort sao portas pequenas.", "Evita contratos grandes e acoplamento desnecessario."),
        ("Dependency Inversion", "Modulos de alto nivel devem depender de abstracoes, nao de detalhes.", "Casos de uso recebem AuthPort, WalletRepository, EventRepository e BetRepository no construtor.", "SQL, JWT e Oracle ficam fora da regra de negocio."),
    ]
    for name, explanation, evidence, effect in solid_rows:
        doc.add_heading(name, level=2)
        doc.add_paragraph(f"Explicacao tecnica: {explanation}")
        doc.add_paragraph(f"Trecho aplicado: {evidence}")
        doc.add_paragraph(f"Efeito: {effect}")
    add_code(doc, """
        export class FinishEventUseCase {
            constructor(
                private readonly auth: AuthPort,
                private readonly events: EventRepository,
                private readonly bets: BetRepository,
                private readonly wallets: WalletRepository,
                private readonly payoutStrategy: PayoutStrategy
            ) {}
        }
    """)

    doc.add_heading("5. Clean Code", level=1)
    doc.add_paragraph("As praticas de Clean Code foram aplicadas principalmente nos fluxos que concentram risco de negocio: carteira, aposta e finalizacao de eventos.")
    add_table(doc, ["Pratica", "Evidencia", "Efeito"], [
        ["Nomes claros", "MoveWalletFundsUseCase, PlaceBetUseCase, FinishEventUseCase, PoolPayoutStrategy.", "O nome descreve a intencao e reduz necessidade de comentarios."],
        ["Funcoes pequenas", "toCents, assertPositiveCents, WalletTransactionFactory.deposit/withdraw/bet/win.", "Cada funcao tem responsabilidade curta e testavel."],
        ["Reducao de duplicacao", "Deposito e saque usam MoveWalletFundsUseCase.", "Evita duas implementacoes quase iguais para movimentacao de saldo."],
        ["Baixo acoplamento", "Casos de uso dependem de portas, nao de OracleDB diretamente.", "Infraestrutura pode mudar com menor impacto."],
        ["Tratamento explicito de erros", "UseCaseResult com ok/fail.", "Fluxos retornam erro controlado em vez de excecoes espalhadas."],
    ], widths=[1900, 4000, 3460])
    add_code(doc, """
        export function toCents(amount: number): number {
            if (!Number.isFinite(amount) || amount <= 0) {
                throw new Error("O valor deve ser maior que zero.");
            }

            return Math.round(amount * 100);
        }
    """)

    doc.add_heading("6. Padroes de Projeto GoF", level=1)
    doc.add_heading("Strategy - comportamento", level=2)
    doc.add_paragraph(
        "Problema: o calculo de premiacao pode mudar. A solucao foi criar PayoutStrategy e uma implementacao PoolPayoutStrategy. Beneficio: extensibilidade. Custo: mais uma abstracao para entender."
    )
    add_code(doc, """
        export interface PayoutStrategy {
            calculatePayouts(winningGuess: string, bets: SettledBet[]): Payout[];
        }

        export class PoolPayoutStrategy implements PayoutStrategy {
            calculatePayouts(winningGuess: string, bets: SettledBet[]): Payout[] {
                const totalPool = bets.reduce((sum, bet) => sum + bet.betAmountInCents, 0);
                // ... distribuicao proporcional aos vencedores
            }
        }
    """)
    doc.add_heading("Factory - criacao", level=2)
    doc.add_paragraph(
        "Problema: transacoes de carteira precisam ser consistentes quanto a tipo e validacao. A Factory padroniza deposito, saque, aposta e ganho. Beneficio: centralizacao da criacao. Custo: chamada indireta em vez de objeto literal."
    )
    add_code(doc, """
        export class WalletTransactionFactory {
            static deposit(walletId: number, amountInCents: number): WalletTransaction {
                return this.create(walletId, amountInCents, "Deposito");
            }

            static win(walletId: number, amountInCents: number): WalletTransaction {
                return this.create(walletId, amountInCents, "Ganho");
            }
        }
    """)
    doc.add_heading("Adapter - estrutura", level=2)
    doc.add_paragraph(
        "Problema: casos de uso nao devem depender diretamente de OracleDB, JWT ou bibliotecas externas. Os adapters traduzem essas APIs para portas internas. Beneficio: substituicao e teste. Custo: mais classes de infraestrutura."
    )
    add_code(doc, """
        export class OracleWalletRepository implements WalletRepository {
            constructor(private readonly connection: OracleDB.Connection) {}

            async findByOwnerId(ownerId: number): Promise<StoredWallet | null> {
                const result = await this.connection.execute(
                    `SELECT id, balance, owner_id FROM wallets WHERE owner_id = :ownerId`,
                    [ownerId],
                    { outFormat: OracleDB.OUT_FORMAT_OBJECT }
                );
                // ... mapeia linha Oracle para StoredWallet
            }
        }
    """)
    add_image(doc, RENDERED / "gof-class-diagram.png", "Figura 2 - Classes dos padroes GoF. Fonte: diagrams/gof-class-diagram.mmd.")

    doc.add_heading("7. Design de API", level=1)
    doc.add_paragraph(
        "O estilo escolhido e REST, pois o sistema ja possui recursos claros como account, wallet, event, bet e moderation. REST tambem combina com Express e com ferramentas de teste e documentacao amplamente conhecidas."
    )
    add_bullets(doc, [
        "Especificacao formal: docs/openapi.yaml em OpenAPI 3.0.3.",
        "Autenticacao: Bearer JWT no header Authorization.",
        "Versionamento: sem versao explicita na URL nesta iteracao; o contrato e versionado pelo campo info.version do OpenAPI. Mudancas incompativeis futuras devem migrar para /v1 ou novo contrato.",
        "Convencao de erro: respostas JSON com propriedade error.",
        "Paginacao: ainda nao aplicada porque os endpoints de lista sao simples; recomendacao futura para GET /event/getEvents e GET /event/getMyEvents."
    ])
    add_table(doc, ["Endpoint", "Responsabilidade"], [
        ["POST /account/signUp", "Cadastro de usuario."],
        ["POST /account/login", "Login e emissao de token."],
        ["POST /account/addFunds", "Deposito em carteira."],
        ["POST /account/withdrawFunds", "Saque de carteira."],
        ["POST /event/addEvent", "Criacao de evento pendente."],
        ["POST /event/bet", "Realizacao de aposta."],
        ["POST /mod/evaluateEvent", "Aprovacao ou reprovacao por moderador."],
        ["POST /mod/finishEvent", "Finalizacao e distribuicao de ganhos."],
    ], widths=[2700, 6660])

    doc.add_heading("8. Diagramas e Modelos", level=1)
    doc.add_paragraph("Os diagramas foram mantidos como codigo em Mermaid na pasta /diagrams, conforme recomendado no enunciado.")
    add_bullets(doc, [
        "diagrams/architecture.mmd: visao geral de componentes.",
        "diagrams/gof-class-diagram.mmd: classes onde GoF foi aplicado.",
        "diagrams/place-bet-sequence.mmd: fluxo central de realizacao de aposta."
    ])
    add_image(doc, RENDERED / "place-bet-sequence.png", "Figura 3 - Sequencia do fluxo de aposta. Fonte: diagrams/place-bet-sequence.mmd.")

    doc.add_heading("9. Conclusoes", level=1)
    doc.add_paragraph(
        "A refatoracao arquitetural tornou o projeto mais defensavel: os fluxos centrais passaram a demonstrar separacao de responsabilidades, inversao de dependencias, padroes GoF aplicados a problemas reais e testes automatizados de dominio."
    )
    doc.add_paragraph(
        "A pratica mais efetiva foi separar regra de negocio de infraestrutura. Isso permitiu testar premiacao e transacoes de carteira sem depender de OracleDB ou servidor Express."
    )
    doc.add_heading("Limitacoes e proximas melhorias", level=2)
    add_bullets(doc, [
        "Usuarios legados com senha antiga devem passar por migracao ou troca obrigatoria de senha.",
        "Validacao de entrada deveria ser centralizada com biblioteca como Zod, Joi ou class-validator.",
        "Nem todos os servicos legados foram migrados integralmente para casos de uso.",
        "Faltam testes de integracao para rotas Express e persistencia Oracle.",
        "Os endpoints de listagem ainda nao possuem paginacao.",
        "Mensagens e arquivos antigos possuem alguns problemas de encoding em acentos.",
        "O livro-texto exato da disciplina e os nomes dos integrantes devem ser preenchidos antes da entrega."
    ])

    doc.add_heading("10. Referencias Bibliograficas", level=1)
    references = [
        "GAMMA, Erich; HELM, Richard; JOHNSON, Ralph; VLISSIDES, John. Design Patterns: Elements of Reusable Object-Oriented Software. Reading: Addison-Wesley, 1994.",
        "MARTIN, Robert C. Clean Code: A Handbook of Agile Software Craftsmanship. Upper Saddle River: Prentice Hall, 2008.",
        "MARTIN, Robert C. Clean Architecture: A Craftsman's Guide to Software Structure and Design. Boston: Prentice Hall, 2017.",
        "FOWLER, Martin. Patterns of Enterprise Application Architecture. Boston: Addison-Wesley, 2002.",
        "INTERNATIONAL ORGANIZATION FOR STANDARDIZATION. ISO/IEC 25010:2023: Systems and software engineering - Systems and software Quality Requirements and Evaluation (SQuaRE) - Product quality model. Geneva: ISO, 2023.",
        "NYGARD, Michael. Documenting Architecture Decisions. 2011. Disponivel em: https://cognitect.com/blog/2011/11/15/documenting-architecture-decisions. Acesso em: 10 jun. 2026.",
        "LIVRO-TEXTO DA DISCIPLINA. Referencia a completar conforme bibliografia indicada pelo professor."
    ]
    for ref in references:
        doc.add_paragraph(ref)

    doc.add_heading("Anexos", level=1)
    add_bullets(doc, [
        "Repositorio Git: https://github.com/KachanChK/TCG.Bet",
        "Pasta de ADRs: /adrs",
        "Pasta de diagramas: /diagrams",
        "Contrato de API: /docs/openapi.yaml",
        "README de instalacao e execucao: /README.md",
        "Testes automatizados: /tests",
        "Comandos verificados: npm.cmd run build e npm.cmd test"
    ])

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
